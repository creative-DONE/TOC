import pytest
import os
import sys
from pathlib import Path
from openpyxl import load_workbook

backend_dir = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(backend_dir))

from app.db.database import SessionLocal
from app.models.factory_models import Machine
from app.models.order_models import Order, OrderBatch
from app.models.schedule_models import ProductionSchedule
from app.schemas.schemas import MatrixEditRequest
from app.services.matrix_service import get_planning_matrix_data, execute_matrix_action, extract_numeric_order_id
from app.services.report_service import generate_production_excel_report
from app.core.batch_splitter import split_order_into_capacity_allocations

@pytest.fixture(scope="module")
def db():
    session = SessionLocal()
    yield session
    session.close()

def test_get_planning_matrix_structure(db):
    """Verifies that planning matrix returns dynamic machines, dynamic bottleneck, sticky order data, summaries, and work cells."""
    matrix = get_planning_matrix_data(db, horizon_days=7)
    
    assert "machines" in matrix
    assert "orders" in matrix
    assert "active_constraint" in matrix
    assert "summary" in matrix
    assert "machine_load_summary" in matrix
    assert "day_summary" in matrix

    machines = matrix["machines"]
    assert len(machines) > 0

    for m in machines:
        assert "code" in m
        assert "capacity_kg" in m
        assert "current_load_kg" in m
        assert "utilization_pct" in m
        assert "processing_time_hours" in m
        assert "working_hours_per_day" in m
        assert m["capacity_kg"] > 0
        assert m["processing_time_hours"] > 0

    # Active constraint must be identified dynamically
    active_constraint = matrix["active_constraint"]
    assert "resource_name" in active_constraint
    assert "utilization_pct" in active_constraint

    orders = matrix["orders"]
    assert len(orders) > 0

    for o in orders:
        assert "order_id" in o
        assert "order_number" in o
        assert "short_order_number" in o
        assert "quantity_kg" in o
        assert "planned_day" in o
        assert "machine_cells" in o

        for m in machines:
            m_key = str(m["id"])
            assert m_key in o["machine_cells"]
            cell = o["machine_cells"][m_key]
            assert "assigned" in cell
            if cell["assigned"]:
                assert "cell_display" in cell

def test_matrix_numeric_order_sorting(db):
    """Verifies that orders sort in ascending numeric order (101, 102, 103...), NOT string order."""
    matrix = get_planning_matrix_data(db, horizon_days=7)
    orders = matrix["orders"]
    numeric_ids = [extract_numeric_order_id(o.get("raw_order_number", o["order_number"])) for o in orders]
    # Check that numeric_ids are monotonically non-decreasing
    assert numeric_ids == sorted(numeric_ids), f"Order sequence {numeric_ids} is not strictly sorted by numeric key."

def test_capacity_splitting_algorithm():
    """Verifies CEILING(order_qty / machine_cap) auto-splitting and sum invariants."""
    # Case 1: 8000 kg on 4000 kg capacity -> 2 batches of 4000 kg
    res_8000 = split_order_into_capacity_allocations(8000.0, 4000.0)
    assert len(res_8000) == 2
    assert res_8000 == [4000.0, 4000.0]
    assert sum(res_8000) == 8000.0

    # Case 2: 9500 kg on 4000 kg capacity -> 3 batches: 4000, 4000, 1500
    res_9500 = split_order_into_capacity_allocations(9500.0, 4000.0)
    assert len(res_9500) == 3
    assert res_9500 == [4000.0, 4000.0, 1500.0]
    assert sum(res_9500) == 9500.0

    # Case 3: Fits single batch
    res_300 = split_order_into_capacity_allocations(300.0, 500.0)
    assert len(res_300) == 1
    assert res_300 == [300.0]

def test_matrix_auto_splitting_on_assignment(db):
    """Tests that assigning an oversized order to a machine automatically creates valid sequential allocations."""
    # Find smallest capacity machine
    matrix = get_planning_matrix_data(db)
    smallest_mach = min(matrix["machines"], key=lambda m: m["capacity_kg"])
    cap = smallest_mach["capacity_kg"]

    # Target order with 2.5x capacity
    target_order = matrix["orders"][0]
    ord_id = target_order["order_id"]
    test_qty = cap * 2.5  # e.g., 500 * 2.5 = 1250 kg

    res = execute_matrix_action(db, MatrixEditRequest(
        action="ASSIGN",
        order_id=ord_id,
        target_machine_id=smallest_mach["id"],
        quantity_kg=test_qty
    ))

    assert res["success"] is True

    # Verify slots created in DB for this order
    slots = db.query(ProductionSchedule).filter(ProductionSchedule.order_id == ord_id).all()
    assert len(slots) == 3  # ceil(2.5) = 3

    # Verify every slot weight <= machine capacity and sum == test_qty
    allocated_weights = [s.batch.batch_quantity_kg if s.batch else s.order.quantity_kg for s in slots]
    for w in allocated_weights:
        assert w <= cap, f"Slot weight {w} exceeds machine capacity {cap}"
    assert round(sum(allocated_weights), 1) == round(test_qty, 1)

def test_matrix_pre_allocation_validation(db):
    """Tests that invalid inputs (e.g. quantity <= 0) are strictly rejected with validation errors."""
    res = execute_matrix_action(db, MatrixEditRequest(
        action="ASSIGN",
        order_id=1,
        target_machine_id=1,
        quantity_kg=-50.0
    ))
    assert res["success"] is False
    assert "VALIDATION ERROR" in res["error"]

def test_matrix_toggle_lock(db):
    """Tests locking and unlocking an order from the matrix."""
    matrix = get_planning_matrix_data(db)
    target_order = matrix["orders"][0]
    ord_id = target_order["order_id"]

    res1 = execute_matrix_action(db, MatrixEditRequest(action="TOGGLE_LOCK", order_id=ord_id))
    assert res1["success"] is True
    assert "is_locked" in res1

    res2 = execute_matrix_action(db, MatrixEditRequest(action="TOGGLE_LOCK", order_id=ord_id))
    assert res2["success"] is True
    assert res2["is_locked"] != res1["is_locked"]

def test_matrix_update_row_and_cascade(db):
    """Tests editing row parameters with full plant recalculation."""
    matrix = get_planning_matrix_data(db)
    target_order = matrix["orders"][0]
    ord_id = target_order["order_id"]

    res = execute_matrix_action(db, MatrixEditRequest(
        action="UPDATE_ROW",
        order_id=ord_id,
        quantity_kg=target_order["quantity_kg"],
        planned_day=2
    ))

    assert res["success"] is True
    assert "SCHEDULE RE-OPTIMIZED" in res["message"]
    assert "diff" in res

def test_matrix_move_and_compact(db):
    """Tests moving an order to another machine without collision and re-sequencing."""
    matrix = get_planning_matrix_data(db)
    assigned_order = None
    source_m_id = None
    for o in matrix["orders"]:
        for m_id_str, cell in o["machine_cells"].items():
            if cell["assigned"]:
                assigned_order = o
                source_m_id = int(m_id_str)
                break
        if assigned_order:
            break

    assert assigned_order is not None

    target_mach = next((m for m in matrix["machines"] if m["id"] != source_m_id), None)
    assert target_mach is not None

    res = execute_matrix_action(db, MatrixEditRequest(
        action="MOVE",
        order_id=assigned_order["order_id"],
        machine_id=source_m_id,
        target_machine_id=target_mach["id"]
    ))

    assert res["success"] is True
    assert "SCHEDULE RE-OPTIMIZED" in res["message"]

def test_matrix_sequential_days(db):
    """Verifies that planned days sequence compactly as Day 1, Day 2, Day 3... without artificial gaps."""
    matrix = get_planning_matrix_data(db, horizon_days=7)
    day_summary = matrix["day_summary"]
    assert len(day_summary) == 7
    days = [d["day"] for d in day_summary]
    assert days == [1, 2, 3, 4, 5, 6, 7]

def test_excel_planning_matrix_export(db):
    """Verifies that generated Excel file includes Planning Matrix sheet with WORK header and allocated quantity formatting."""
    schedules = db.query(ProductionSchedule).all()
    filepath = generate_production_excel_report(schedules, report_title="Production Planning Matrix", db=db)

    assert os.path.exists(filepath)
    wb = load_workbook(filepath)

    assert "Planning Matrix" in wb.sheetnames
    assert "Detailed Slots" in wb.sheetnames

    ws = wb["Planning Matrix"]
    assert "PRIME TEXTILES" in str(ws["A1"].value)
    assert ws["A3"].value == "ORDER SPECIFICATIONS"
    assert ws["E3"].value == "WORK"

    assert ws["A4"].value == "Order Number"
    assert ws["B4"].value == "Quantity (kg)"
    assert ws["C4"].value == "Due Date"
    assert ws["D4"].value == "Planned Day"

    found_assigned_cell = False
    for r in range(5, ws.max_row + 1):
        for c in range(5, ws.max_column + 1):
            val = ws.cell(row=r, column=c).value
            if val is not None and str(val).strip() != "":
                found_assigned_cell = True
                assert "kg" in str(val) or str(val).isdigit() or "-" in str(val)
                break
        if found_assigned_cell:
            break

    assert found_assigned_cell is True

def test_matrix_machine_batch_processing_time(db):
    """Verifies that planning matrix machine headers include editable processing_time_hours and dynamic updates cascade."""
    matrix = get_planning_matrix_data(db, horizon_days=7)
    machines = matrix["machines"]

    m1 = next((m for m in machines if m["code"] == "JET-M1"), None)
    assert m1 is not None
    assert "processing_time_hours" in m1
    assert "working_hours_per_day" in m1
    assert m1["processing_time_hours"] > 0
    assert m1["working_hours_per_day"] == 8.0

    # Test dynamic update: change M1 processing time from current to 4.5 hours
    db_m1 = db.query(Machine).filter(Machine.code == "JET-M1").first()
    original_proc_time = db_m1.processing_time_hours
    try:
        db_m1.processing_time_hours = 4.5
        db.commit()

        updated_matrix = get_planning_matrix_data(db, horizon_days=7)
        updated_m1 = next((m for m in updated_matrix["machines"] if m["code"] == "JET-M1"), None)
        assert updated_m1["processing_time_hours"] == 4.5
        # With longer batch time, nominal capacity should decrease
        assert updated_m1["nominal_capacity_kg"] < m1["nominal_capacity_kg"]
    finally:
        db_m1.processing_time_hours = original_proc_time
        db.commit()

def test_order_deletion_rebalances_flexible_machines(db):
    """
    Verifies that when an order is deleted:
    1. The central scheduling engine immediately re-evaluates remaining orders.
    2. Available machine capacity is re-evaluated across the fleet.
    3. The notification message strictly follows the required format:
       'Schedule recalculated — X orders reassigned to available capacity.' or
       'Schedule recalculated — no reassignment required.'
    """
    from datetime import datetime, timedelta
    now = datetime.utcnow()
    temp_order = Order(
        order_number=f"ORD-DEL-{int(now.timestamp()) % 10000}",
        customer_id=1,
        cloth_type="Cotton 100% Greige Knit",
        colour_name="Navy Blue",
        colour_code="NAVY",
        quantity_kg=400.0,
        due_date=now + timedelta(days=5),
        order_date=now,
        priority="HIGH",
        status="PENDING"
    )
    db.add(temp_order)
    db.commit()
    db.refresh(temp_order)

    m6 = db.query(Machine).filter(Machine.code == "STENT-M6").first()
    res_assign = execute_matrix_action(db, MatrixEditRequest(
        action="ASSIGN",
        order_id=temp_order.id,
        target_machine_id=m6.id,
        quantity_kg=400.0,
        planned_day=2
    ))
    assert res_assign["success"] is True

    res_del = execute_matrix_action(db, MatrixEditRequest(
        action="DELETE_ORDER",
        order_id=temp_order.id
    ))
    assert res_del["success"] is True
    assert "Schedule recalculated" in res_del["message"]
    reassigned_count = res_del.get("diff", {}).get("reassigned_count", 0)
    if reassigned_count > 0:
        expected_msg = f"Schedule recalculated — {reassigned_count} order{'s' if reassigned_count > 1 else ''} reassigned to available capacity."
        assert res_del["message"] == expected_msg
    else:
        assert res_del["message"] == "Schedule recalculated — no reassignment required."

def test_locked_orders_protected_from_automatic_rebalance(db):
    """
    Verifies that locked orders (is_locked=True) are strictly protected and never
    moved by automatic schedule rebalancing when other orders are modified/deleted.
    """
    matrix = get_planning_matrix_data(db)
    orders = matrix["orders"]
    assert len(orders) > 1

    target = orders[0]
    target_id = target["order_id"]
    orig_assigned_m = target.get("assigned_machine_id")

    res_lock = execute_matrix_action(db, MatrixEditRequest(
        action="TOGGLE_LOCK",
        order_id=target_id
    ))
    assert res_lock["success"] is True
    if not res_lock.get("is_locked"):
        res_lock = execute_matrix_action(db, MatrixEditRequest(
            action="TOGGLE_LOCK",
            order_id=target_id
        ))
    assert res_lock.get("is_locked") is True

    ord_obj = db.query(Order).filter(Order.id == target_id).first()
    assert ord_obj.is_locked is True

    res_reopt = execute_matrix_action(db, MatrixEditRequest(
        action="REOPTIMIZE"
    ))
    assert res_reopt["success"] is True

    db.refresh(ord_obj)
    assert ord_obj.is_locked is True
    if orig_assigned_m:
        assert ord_obj.assigned_machine_id == orig_assigned_m

    res_unlock = execute_matrix_action(db, MatrixEditRequest(
        action="TOGGLE_LOCK",
        order_id=target_id
    ))
    assert res_unlock["success"] is True
    assert res_unlock.get("is_locked") is False

def test_matrix_cell_order_number_and_quantity_separated(db):
    """
    Verifies that machine cells in the planning matrix keep order_number and quantity_kg
    distinct, format cell_display with line separation, and never concatenate them into
    an ambiguous continuous string like '110250 kg'.
    """
    matrix = get_planning_matrix_data(db)
    orders = matrix["orders"]
    assert len(orders) > 0

    assigned_count = 0
    for o in orders:
        for m_id_str, cell in o["machine_cells"].items():
            if cell["assigned"]:
                assigned_count += 1
                # Must contain distinct fields
                assert "order_number" in cell
                assert "quantity_kg" in cell
                assert "cell_display" in cell

                # Order number must be prefixed (e.g. ORD-101)
                assert cell["order_number"].startswith("ORD-")
                # cell_display must clearly separate order number and quantity with a newline
                assert "\n" in cell["cell_display"]
                parts = cell["cell_display"].split("\n")
                assert len(parts) == 2
                assert parts[0].startswith("ORD-")
                assert "kg" in parts[1]
                # Cannot be concatenated like "110250 kg"
                assert not parts[0].replace("ORD-", "").endswith(str(int(cell["quantity_kg"])))

    assert assigned_count > 0, "Expected at least one assigned machine cell in the planning matrix."

