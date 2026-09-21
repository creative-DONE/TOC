import pytest
import os
import sys
from pathlib import Path
from datetime import datetime, timedelta
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from openpyxl import load_workbook

backend_dir = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(backend_dir))

from app.db.database import Base
from app.models.factory_models import (
    Machine, MachineMaintenance, FactoryUtility, Employee, MachineStatus, MachineType
)
from app.models.order_models import Order, Customer, OrderStatus, OrderBatch
from app.models.schedule_models import ProductionSchedule
from app.schemas.schemas import MatrixEditRequest
from app.services.scheduler_service import (
    optimize_factory_schedule, add_production_time_over_shifts
)
from app.services.order_calculator_service import calculate_order_time_estimate
from app.services.matrix_service import execute_matrix_action, get_planning_matrix_data
from app.services.report_service import generate_production_excel_report


@pytest.fixture
def test_db():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(bind=engine)
    Session = sessionmaker(bind=engine)
    session = Session()

    now = datetime(2026, 9, 21, 8, 0, 0)

    # Customer
    cust = Customer(id=1, code="CUST-001", name="Prime Garments", priority_tier="TIER_1", importance_weight=1.0)
    session.add(cust)

    # Fleet with manual time configurations:
    # JET-M1: Cap 500, Proc 3.0h, Load 0.5h, Unload 0.5h, Clean 1.0h
    # JET-M2: Cap 800, Proc 4.0h, Load 0.6h, Unload 0.6h, Clean 1.2h
    # SOFT-M3: Cap 400, Proc 2.5h, Load 0.4h, Unload 0.4h, Clean 0.8h
    # SOFT-M4: Cap 600, Proc 3.0h, Load 0.5h, Unload 0.5h, Clean 1.0h
    # JIG-M5: Cap 1000, Proc 5.0h, Load 1.0h, Unload 1.0h, Clean 1.5h
    # STENT-M6: Cap 1500, Proc 2.0h, Load 0.5h, Unload 0.5h, Clean 1.0h
    machines = [
        Machine(
            id=1, code="JET-M1", name="Jet Dyeing M1", machine_type=MachineType.JET_DYEING.value,
            capacity_kg=500.0, min_batch_kg=50.0, max_batch_kg=500.0, processing_speed=1.0, efficiency=0.92,
            processing_time_hours=3.0, loading_time_hours=0.5, unloading_time_hours=0.5, cleaning_time_hours=1.0,
            working_hours_per_day=8.0, status=MachineStatus.AVAILABLE.value, power_kw=45.0, water_m3_hr=3.5, steam_kg_hr=600.0,
            compatible_cloth_types="Cotton,Polyester,Poly-Cotton Blend,Rayon Viscose", compatible_colours="ALL"
        ),
        Machine(
            id=2, code="JET-M2", name="Jet Dyeing M2", machine_type=MachineType.JET_DYEING.value,
            capacity_kg=800.0, min_batch_kg=100.0, max_batch_kg=800.0, processing_speed=1.1, efficiency=0.94,
            processing_time_hours=4.0, loading_time_hours=0.6, unloading_time_hours=0.6, cleaning_time_hours=1.2,
            working_hours_per_day=8.0, status=MachineStatus.AVAILABLE.value, power_kw=65.0, water_m3_hr=5.2, steam_kg_hr=900.0,
            compatible_cloth_types="Cotton,Polyester,Poly-Cotton Blend", compatible_colours="ALL"
        ),
        Machine(
            id=3, code="SOFT-M3", name="Soft Flow Vessel M3", machine_type=MachineType.SOFT_FLOW.value,
            capacity_kg=400.0, min_batch_kg=50.0, max_batch_kg=400.0, processing_speed=0.95, efficiency=0.90,
            processing_time_hours=2.5, loading_time_hours=0.4, unloading_time_hours=0.4, cleaning_time_hours=0.8,
            working_hours_per_day=8.0, status=MachineStatus.AVAILABLE.value, power_kw=35.0, water_m3_hr=2.8, steam_kg_hr=500.0,
            compatible_cloth_types="Cotton,Organic Cotton,Rayon Viscose,Modal", compatible_colours="ALL"
        ),
        Machine(
            id=4, code="SOFT-M4", name="Soft Flow Vessel M4", machine_type=MachineType.SOFT_FLOW.value,
            capacity_kg=600.0, min_batch_kg=100.0, max_batch_kg=600.0, processing_speed=1.0, efficiency=0.91,
            processing_time_hours=3.0, loading_time_hours=0.5, unloading_time_hours=0.5, cleaning_time_hours=1.0,
            working_hours_per_day=8.0, status=MachineStatus.AVAILABLE.value, power_kw=48.0, water_m3_hr=4.0, steam_kg_hr=700.0,
            compatible_cloth_types="Cotton,Poly-Cotton Blend,Rayon Viscose", compatible_colours="ALL"
        ),
        Machine(
            id=5, code="JIG-M5", name="Jigger Dyeing J5", machine_type=MachineType.JIGGER.value,
            capacity_kg=1000.0, min_batch_kg=200.0, max_batch_kg=1000.0, processing_speed=0.85, efficiency=0.88,
            processing_time_hours=5.0, loading_time_hours=1.0, unloading_time_hours=1.0, cleaning_time_hours=1.5,
            working_hours_per_day=8.0, status=MachineStatus.AVAILABLE.value, power_kw=55.0, water_m3_hr=4.5, steam_kg_hr=750.0,
            compatible_cloth_types="Cotton,Woven Twill,Denim,Linen Blend", compatible_colours="ALL"
        ),
        Machine(
            id=6, code="STENT-M6", name="Stenter Line S6", machine_type=MachineType.STENTER_FINISHING.value,
            capacity_kg=1500.0, min_batch_kg=100.0, max_batch_kg=1500.0, processing_speed=1.3, efficiency=0.95,
            processing_time_hours=2.0, loading_time_hours=0.5, unloading_time_hours=0.5, cleaning_time_hours=1.0,
            working_hours_per_day=8.0, status=MachineStatus.AVAILABLE.value, power_kw=90.0, water_m3_hr=1.5, steam_kg_hr=1200.0,
            compatible_cloth_types="Cotton,Polyester,Poly-Cotton Blend,Rayon Viscose,Organic Cotton", compatible_colours="ALL"
        ),
    ]
    session.add_all(machines)

    # Factory Utilities
    session.add_all([
        FactoryUtility(name="Electricity", capacity_per_hour=500.0, current_load_per_hour=150.0, unit="kW"),
        FactoryUtility(name="Water", capacity_per_hour=30.0, current_load_per_hour=12.0, unit="m3/hr"),
        FactoryUtility(name="Steam", capacity_per_hour=6000.0, current_load_per_hour=2000.0, unit="kg/hr"),
    ])

    # Employee
    session.add(Employee(id=1, code="EMP-01", name="Senior Dyer", role="OPERATOR", shift="SHIFT_A", on_leave=False))

    session.commit()
    yield session
    session.close()


def test_test1_multibatch_order_zero_cleaning_between_batches(test_db):
    """
    TEST 1: Multi-batch order has 0h cleaning between Batch 1 and Batch 2.
    An 800 kg order on M1 (capacity 500 kg) splits into 2 batches (500 kg and 300 kg).
    Batch 1 has cleaning (if color change) or 0h; Batch 2 MUST have exactly 0h cleaning.
    """
    ref_now = datetime(2026, 9, 21, 8, 0, 0)
    order = Order(
        id=101, order_number="ORD-MB-1", customer_id=1, cloth_type="Cotton",
        colour_name="Navy Blue", colour_code="NAVY", quantity_kg=800.0,
        due_date=ref_now + timedelta(days=3), order_date=ref_now, status="PENDING"
    )
    test_db.add(order)
    test_db.commit()

    res = execute_matrix_action(test_db, MatrixEditRequest(
        action="ASSIGN", order_id=101, target_machine_id=1, quantity_kg=800.0
    ), reference_now=ref_now)
    assert res["success"] is True

    slots = test_db.query(ProductionSchedule).filter(ProductionSchedule.order_id == 101).order_by(ProductionSchedule.planned_start.asc()).all()
    assert len(slots) == 2

    # Batch 2 must have exactly 0.0 cleaning minutes
    assert slots[1].cleaning_min == 0.0
    assert slots[1].changeover_min == 0.0
    assert slots[1].loading_min == 30.0   # 0.5h
    assert slots[1].base_processing_min == 180.0  # 3.0h
    assert slots[1].unloading_min == 30.0 # 0.5h


def test_test2_consecutive_orders_same_color_zero_cleaning(test_db):
    """
    TEST 2: Consecutive orders with the same color on the same machine have 0h cleaning.
    """
    ref_now = datetime(2026, 9, 21, 8, 0, 0)
    o1 = Order(
        id=102, order_number="ORD-COL-1", customer_id=1, cloth_type="Cotton",
        colour_name="Royal Blue", colour_code="ROYAL_BLUE", quantity_kg=400.0,
        due_date=ref_now + timedelta(days=2), order_date=ref_now, status="PENDING"
    )
    o2 = Order(
        id=103, order_number="ORD-COL-2", customer_id=1, cloth_type="Cotton",
        colour_name="Royal Blue", colour_code="ROYAL_BLUE", quantity_kg=400.0,
        due_date=ref_now + timedelta(days=3), order_date=ref_now, status="PENDING"
    )
    test_db.add_all([o1, o2])
    test_db.commit()

    optimize_factory_schedule(test_db, reference_now=ref_now, horizon_days=7)

    # Check slots for o1 and o2 if assigned to same machine
    s1 = test_db.query(ProductionSchedule).filter(ProductionSchedule.order_id == 102).first()
    s2 = test_db.query(ProductionSchedule).filter(ProductionSchedule.order_id == 103).first()
    assert s1 is not None and s2 is not None

    # If scheduled consecutively on the same machine, s2 must have 0.0 cleaning_min
    if s1.machine_id == s2.machine_id and s2.planned_start >= s1.planned_end:
        assert s2.cleaning_min == 0.0


def test_test3_orders_different_colors_configured_cleaning_added(test_db):
    """
    TEST 3: Orders with different colors on the same machine have configured cleaning time added.
    """
    ref_now = datetime(2026, 9, 21, 8, 0, 0)
    # Target M3 (cleaning_time_hours = 0.8h -> 48 min)
    m3 = test_db.query(Machine).filter(Machine.id == 3).first()
    assert m3.cleaning_time_hours == 0.8

    o1 = Order(
        id=104, order_number="ORD-DIFF-1", customer_id=1, cloth_type="Cotton",
        colour_name="Pure White", colour_code="WHITE", quantity_kg=300.0,
        due_date=ref_now + timedelta(days=2), order_date=ref_now, status="PENDING"
    )
    o2 = Order(
        id=105, order_number="ORD-DIFF-2", customer_id=1, cloth_type="Cotton",
        colour_name="Jet Black", colour_code="BLACK", quantity_kg=300.0,
        due_date=ref_now + timedelta(days=3), order_date=ref_now, status="PENDING"
    )
    test_db.add_all([o1, o2])
    test_db.commit()

    # Assign both manually to M3 sequentially
    res1 = execute_matrix_action(test_db, MatrixEditRequest(
        action="ASSIGN", order_id=104, target_machine_id=3, quantity_kg=300.0, planned_day=1
    ), reference_now=ref_now)
    assert res1["success"] is True

    # Order 2 has different color (BLACK vs previous WHITE)
    res2 = execute_matrix_action(test_db, MatrixEditRequest(
        action="ASSIGN", order_id=105, target_machine_id=3, quantity_kg=300.0, planned_day=2
    ), reference_now=ref_now)
    assert res2["success"] is True

    s2 = test_db.query(ProductionSchedule).filter(ProductionSchedule.order_id == 105).first()
    assert s2 is not None
    assert s2.cleaning_min == pytest.approx(48.0, 0.1)  # 0.8h * 60 = 48 min


def test_test4_order_three_batches_zero_cleaning_across_all_batches(test_db):
    """
    TEST 4: Order split into 3 batches has 0h cleaning between all 3 batches.
    """
    ref_now = datetime(2026, 9, 21, 8, 0, 0)
    # Order of 1200 kg on M1 (capacity 500 kg) -> 3 batches: 500kg, 500kg, 200kg
    order = Order(
        id=106, order_number="ORD-3B", customer_id=1, cloth_type="Cotton",
        colour_name="Emerald Green", colour_code="GREEN", quantity_kg=1200.0,
        due_date=ref_now + timedelta(days=4), order_date=ref_now, status="PENDING"
    )
    test_db.add(order)
    test_db.commit()

    res = execute_matrix_action(test_db, MatrixEditRequest(
        action="ASSIGN", order_id=106, target_machine_id=1, quantity_kg=1200.0
    ), reference_now=ref_now)
    assert res["success"] is True

    slots = test_db.query(ProductionSchedule).filter(ProductionSchedule.order_id == 106).order_by(ProductionSchedule.planned_start.asc()).all()
    assert len(slots) == 3

    # Batch 2 and Batch 3 must have 0.0 cleaning time
    assert slots[1].cleaning_min == 0.0
    assert slots[2].cleaning_min == 0.0


def test_test5_slot_moves_around_scheduled_maintenance(test_db):
    """
    TEST 5: Production slot moves around a scheduled maintenance window.
    """
    ref_now = datetime(2026, 9, 21, 8, 0, 0)
    # Add maintenance on M2 from 09:00 to 12:00 on Day 1
    maint = MachineMaintenance(
        machine_id=2, title="Pump overhaul", maintenance_type="PREVENTIVE",
        start_time=ref_now + timedelta(hours=1),
        end_time=ref_now + timedelta(hours=4),
        status="SCHEDULED"
    )
    test_db.add(maint)
    test_db.commit()

    # Slot needing 4.0h total duration on M2 starting at 08:00
    # Between 08:00 and 09:00: 1.0h of work
    # 09:00 to 12:00: MAINTENANCE -> paused
    # 12:00 to 15:00: 3.0h remaining work completed at 15:00
    s_start, s_end = add_production_time_over_shifts(
        start_time=ref_now,
        duration_minutes=240.0,  # 4 hours
        working_hours_per_day=8.0,
        maintenances=[maint]
    )
    assert s_start == ref_now
    assert s_end == ref_now + timedelta(hours=7)  # 8:00 + 1h work + 3h maint + 3h work = 15:00


def test_test6_production_exceeding_8h_continues_next_working_day(test_db):
    """
    TEST 6: Production exceeding 8 hours continues on next working day at 08:00 AM.
    Shift: 08:00 to 16:00 (8 hours).
    A job requiring 10 hours starting at 10:00 AM:
    Day 1: 10:00 AM to 04:00 PM (6 hours)
    Day 2: 08:00 AM to 12:00 PM (4 hours remaining) -> Completes Day 2 at 12:00 PM.
    """
    start_time = datetime(2026, 9, 21, 10, 0, 0)
    duration_min = 600.0  # 10 hours

    s_start, s_end = add_production_time_over_shifts(
        start_time=start_time,
        duration_minutes=duration_min,
        working_hours_per_day=8.0,
        maintenances=[]
    )
    assert s_start == start_time
    assert s_end == datetime(2026, 9, 22, 12, 0, 0)


def test_test7_deleting_order_triggers_rebalance_including_m6(test_db):
    """
    TEST 7: Deleting orders from a machine frees capacity and re-evaluates flexible orders.
    """
    ref_now = datetime(2026, 9, 21, 8, 0, 0)
    # Add an order assigned to M6
    o_m6 = Order(
        id=107, order_number="ORD-M6-DEL", customer_id=1, cloth_type="Cotton",
        colour_name="Grey", colour_code="GREY", quantity_kg=1000.0,
        due_date=ref_now + timedelta(days=2), order_date=ref_now, status="PENDING"
    )
    # Add a flexible order that can go to M6
    o_flex = Order(
        id=108, order_number="ORD-FLEX-1", customer_id=1, cloth_type="Cotton",
        colour_name="Grey", colour_code="GREY", quantity_kg=800.0,
        due_date=ref_now + timedelta(days=2), order_date=ref_now, status="PENDING"
    )
    test_db.add_all([o_m6, o_flex])
    test_db.commit()

    # Assign o_m6 to M6
    execute_matrix_action(test_db, MatrixEditRequest(
        action="ASSIGN", order_id=107, target_machine_id=6, quantity_kg=1000.0
    ), reference_now=ref_now)

    # Delete o_m6 -> M6 becomes available and optimizer should place o_flex on compatible machine
    res_del = execute_matrix_action(test_db, MatrixEditRequest(
        action="DELETE_ORDER", order_id=107
    ), reference_now=ref_now)
    assert res_del["success"] is True

    # Verify o_flex has scheduled slots
    slots_flex = test_db.query(ProductionSchedule).filter(ProductionSchedule.order_id == 108).all()
    assert len(slots_flex) > 0


def test_test8_excel_export_includes_utilization_and_time_columns(test_db):
    """
    TEST 8: Excel export includes standardized Utilization % and comprehensive time columns:
    Loading (min), Proc Time (min), Unloading (min), Cleaning (min), Total Time (min).
    """
    ref_now = datetime(2026, 9, 21, 8, 0, 0)
    order = Order(
        id=109, order_number="ORD-EXCEL-TEST", customer_id=1, cloth_type="Cotton",
        colour_name="Red", colour_code="RED", quantity_kg=400.0,
        due_date=ref_now + timedelta(days=2), order_date=ref_now, status="PENDING"
    )
    test_db.add(order)
    test_db.commit()

    execute_matrix_action(test_db, MatrixEditRequest(
        action="ASSIGN", order_id=109, target_machine_id=1, quantity_kg=400.0
    ), reference_now=ref_now)

    schedules = test_db.query(ProductionSchedule).all()
    filepath = generate_production_excel_report(schedules, report_title="Production Planning Matrix", db=test_db)
    assert os.path.exists(filepath)

    wb = load_workbook(filepath)
    assert "Planning Matrix" in wb.sheetnames
    assert "Detailed Slots" in wb.sheetnames

    slots_sheet = wb["Detailed Slots"]
    headers = [cell.value for cell in slots_sheet[3]]
    assert "Loading (min)" in headers
    assert "Proc Time (min)" in headers
    assert "Unloading (min)" in headers
    assert "Cleaning (min)" in headers
    assert "Total Time (min)" in headers

    matrix_sheet = wb["Planning Matrix"]
    sheet_text = " ".join(str(cell.value) for row in matrix_sheet.iter_rows() for cell in row if cell.value)
    assert "FLEET PARAMETERS & REAL PRODUCTION TIME UTILIZATION SUMMARY" in sheet_text
    assert "Utilization %" in sheet_text

    wb.close()
    if os.path.exists(filepath):
        os.remove(filepath)


def test_test9_updating_machine_batch_time_recalculates_schedule(test_db):
    """
    TEST 9: Updating M1 processing_time_hours directly updates duration and utilization.
    """
    ref_now = datetime(2026, 9, 21, 8, 0, 0)
    m1 = test_db.query(Machine).filter(Machine.id == 1).first()
    m1.processing_time_hours = 6.0  # Increase from 3.0 to 6.0 hr
    test_db.commit()

    # Calculator check
    est = calculate_order_time_estimate(
        db=test_db,
        quantity_kg=500.0,
        colour_code="NAVY",
        due_date=ref_now + timedelta(days=3),
        cloth_type="Cotton",
        reference_now=ref_now
    )
    candidates = ([est["recommended_machine"]] if est.get("recommended_machine") else []) + est.get("alternative_machines", [])
    m1_eval = next((c for c in candidates if c and c.get("machine_id") == 1), None)
    assert m1_eval is not None
    # 500kg fits in 1 batch.
    # Total time = loading (0.5h) + proc (6.0h) + unloading (0.5h) + cleaning (1.0h) = 8.0 hours
    assert m1_eval["total_processing_hours"] == pytest.approx(8.0, 0.1)
    assert m1_eval["batches_count"] == 1


def test_test10_changing_loading_unloading_cleaning_times(test_db):
    """
    TEST 10: Changing loading/unloading/cleaning times recalculates schedule and completion time.
    """
    ref_now = datetime(2026, 9, 21, 8, 0, 0)
    m1 = test_db.query(Machine).filter(Machine.id == 1).first()
    m1.loading_time_hours = 1.0     # 60 min
    m1.unloading_time_hours = 1.0   # 60 min
    m1.cleaning_time_hours = 2.0    # 120 min
    m1.processing_time_hours = 3.0  # 180 min
    test_db.commit()

    # Estimate for 1000 kg order -> 2 batches of 500 kg
    # With colour change (NAVY vs preceding WHITE):
    # Batch 1: Loading(1.0) + Proc(3.0) + Unload(1.0) + Clean(2.0) = 7.0 hours
    # Batch 2: Loading(1.0) + Proc(3.0) + Unload(1.0) + Clean(0.0) = 5.0 hours (multi-batch 0h cleaning rule!)
    # Total duration = 12.0 hours
    est_diff = calculate_order_time_estimate(
        db=test_db,
        quantity_kg=1000.0,
        colour_code="NAVY",
        due_date=ref_now + timedelta(days=3),
        cloth_type="Cotton",
        reference_now=ref_now
    )
    candidates_diff = ([est_diff["recommended_machine"]] if est_diff.get("recommended_machine") else []) + est_diff.get("alternative_machines", [])
    m1_diff = next((c for c in candidates_diff if c and c.get("machine_id") == 1), None)
    assert m1_diff is not None
    assert m1_diff["batches_count"] == 2
    assert m1_diff["total_processing_hours"] == pytest.approx(12.0, 0.1)

    # With same colour (WHITE vs preceding WHITE):
    # Batch 1: Loading(1.0) + Proc(3.0) + Unload(1.0) + Clean(0.0) = 5.0 hours
    # Batch 2: Loading(1.0) + Proc(3.0) + Unload(1.0) + Clean(0.0) = 5.0 hours
    # Total duration = 10.0 hours
    est_same = calculate_order_time_estimate(
        db=test_db,
        quantity_kg=1000.0,
        colour_code="WHITE",
        due_date=ref_now + timedelta(days=3),
        cloth_type="Cotton",
        reference_now=ref_now
    )
    candidates_same = ([est_same["recommended_machine"]] if est_same.get("recommended_machine") else []) + est_same.get("alternative_machines", [])
    m1_same = next((c for c in candidates_same if c and c.get("machine_id") == 1), None)
    assert m1_same is not None
    assert m1_same["batches_count"] == 2
    assert m1_same["total_processing_hours"] == pytest.approx(10.0, 0.1)

